from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from decimal import Decimal
from django.core.files.uploadedfile import SimpleUploadedFile
from io import BytesIO
from openpyxl import Workbook, load_workbook

from tests.factories import (
    create_user, create_sucursal, create_product
)
from products.models import Product, StockSucursal, TransferenciaStock, AjusteStock
from cashier.models import Venta, VentaDetalle
from sucursales.models import Sucursal

User = get_user_model()


class ProductStockFlowsTests(TestCase):
    def setUp(self):
        self.suc_a = create_sucursal("Sucursal A")
        self.suc_b = create_sucursal("Sucursal B")
        self.admin = create_user("admin_prod", is_staff=True)
        self.prod = create_product("TP1", "Test Prod", precio_compra=Decimal('1000'), precio_venta=Decimal('2000'))
        # Inicializar stock en sucursal A
        StockSucursal.objects.create(producto=self.prod, sucursal=self.suc_a, cantidad=10)

    def test_transfer_stock_model(self):
        # Simular transferencia manual (modelo)
        before_a = StockSucursal.objects.get(producto=self.prod, sucursal=self.suc_a).cantidad
        ss_b, _ = StockSucursal.objects.get_or_create(producto=self.prod, sucursal=self.suc_b, defaults={'cantidad': 0})
        ss_a = StockSucursal.objects.get(producto=self.prod, sucursal=self.suc_a)
        transfer_qty = 3
        self.assertGreaterEqual(ss_a.cantidad, transfer_qty)
        ss_a.cantidad -= transfer_qty
        ss_a.save(update_fields=['cantidad'])
        ss_b.cantidad += transfer_qty
        ss_b.save(update_fields=['cantidad'])
        TransferenciaStock.objects.create(producto=self.prod, origen=self.suc_a, destino=self.suc_b, cantidad=transfer_qty, usuario=self.admin)
        self.assertEqual(StockSucursal.objects.get(producto=self.prod, sucursal=self.suc_a).cantidad, before_a - transfer_qty)
        self.assertEqual(StockSucursal.objects.get(producto=self.prod, sucursal=self.suc_b).cantidad, transfer_qty)
        self.assertEqual(TransferenciaStock.objects.count(), 1)

    def test_adjust_stock_view_post(self):
        # Ajustar stock vía vista POST
        self.client.force_login(self.admin)
        resp = self.client.post('/products/stock/adjust/', {
            'producto_id': self.prod.id,
            'sucursal_id': self.suc_a.id,
            'delta': 5,
            'motivo': 'Test incremento'
        })
        self.assertEqual(resp.status_code, 200)
        self.prod.refresh_from_db()
        ss = StockSucursal.objects.get(producto=self.prod, sucursal=self.suc_a)
        self.assertEqual(ss.cantidad, 15)  # 10 inicial + 5
        self.assertEqual(AjusteStock.objects.count(), 1)

    def test_adjust_history_filters(self):
        # Crear dos ajustes para probar filtros
        AjusteStock.objects.create(producto=self.prod, sucursal=self.suc_a, cantidad_delta=2)
        AjusteStock.objects.create(producto=self.prod, sucursal=self.suc_b, cantidad_delta=-1)
        resp_all = self.client.get('/products/stock/adjust/history/')
        self.assertEqual(resp_all.status_code, 200)
        html = resp_all.content.decode('utf-8')
        self.assertIn('Historial de ajustes de stock', html)
        # Filtro por sucursal A
        resp_a = self.client.get('/products/stock/adjust/history/', {'sucursal': self.suc_a.id})
        self.assertEqual(resp_a.status_code, 200)
        html_a = resp_a.content.decode('utf-8')
        self.assertIn('+2', html_a)
        # Evitar falso positivo: aseguramos que no aparezca la fila -1 analizando celdas Delta específicas
        self.assertNotIn('>-1<', html_a)

    def test_advanced_reports_ajax_filters_combined(self):
        # Crear ventas para alimentar reporte
        from reports.analytics import compute_analytics
        # Venta artificial: crear venta y detalle directo
        venta = Venta.objects.create(empleado=self.admin, sucursal=self.suc_a, total=Decimal('2000'))
        VentaDetalle.objects.create(venta=venta, producto=self.prod, cantidad=1, precio_unitario=Decimal('2000'))
        fi = (timezone.now() - timezone.timedelta(days=1)).strftime('%Y-%m-%d')
        ff = timezone.now().strftime('%Y-%m-%d')
        self.client.force_login(self.admin)
        resp = self.client.get('/reports/advanced/data/', {
            'fecha_inicio': fi,
            'fecha_fin': ff,
            'cajero': self.admin.id,
            'sucursal': self.suc_a.id,
            'top': 5
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn('top_selling_products', data)
        self.assertTrue(len(data['top_selling_products']) <= 5)
        self.assertIn('ranking_cajeros', data)

    def test_permission_denied_non_staff_reports_json(self):
        non_staff = create_user('empleado1', is_staff=False)
        self.client.force_login(non_staff)
        fi = (timezone.now() - timezone.timedelta(days=1)).strftime('%Y-%m-%d')
        ff = timezone.now().strftime('%Y-%m-%d')
        resp = self.client.get('/reports/advanced/data/', {'fecha_inicio': fi, 'fecha_fin': ff})
        # Debe devolver 302 (redirect) o 403; no 200
        self.assertNotEqual(resp.status_code, 200)

    def test_mass_upload_placeholder(self):
        """Si existe endpoint de carga masiva, simular; si no, marcar skip lógico."""
        # Buscar ruta conocida (ajustar si hay URL específica). Aquí sólo comprobamos que no 404 genérico si existe.
        possible_paths = ['/products/bulk-upload/', '/products/mass-upload/']
        for path in possible_paths:
            resp = self.client.get(path)
            if resp.status_code != 404:
                # Encontrado un endpoint, prueba POST mínimo
                dummy_csv = b"producto_id,nombre,precio_compra,precio_venta\nBULK1,Prod Bulk,500,1200\n"
                file = SimpleUploadedFile('bulk.csv', dummy_csv, content_type='text/csv')
                post_resp = self.client.post(path, {'file': file})
                # No afirmar demasiado: sólo que responde algo procesable
                self.assertIn(post_resp.status_code, (200, 302))
                break

    def test_upload_products_page_and_export_button(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('upload_products'))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode('utf-8')
        self.assertIn('Cargar Productos por Excel', html)
        self.assertIn('Descargar Plantilla', html)
        self.assertIn('Descargar productos actuales', html)

    def test_bulk_assign_products_creates_stock_per_branch(self):
        self.client.force_login(self.admin)
        sucursal = create_sucursal('Sucursal A')
        prod1 = create_product('P1', 'Prod 1', precio_compra=Decimal('1000'), precio_venta=Decimal('1500'))
        prod2 = create_product('P2', 'Prod 2', precio_compra=Decimal('900'), precio_venta=Decimal('1400'))

        resp = self.client.post(reverse('bulk_assign_products'), {
            'products': [prod1.id, prod2.id],
            'sucursal': sucursal.id,
            'cantidad': '5'
        })
        self.assertEqual(resp.status_code, 302)

        ss1 = StockSucursal.objects.get(producto=prod1, sucursal=sucursal)
        ss2 = StockSucursal.objects.get(producto=prod2, sucursal=sucursal)
        self.assertEqual(ss1.cantidad, 5)
        self.assertEqual(ss2.cantidad, 5)
        self.assertEqual(Product.objects.get(id=prod1.id).sucursal_id, sucursal.id)
        self.assertEqual(Product.objects.get(id=prod2.id).sucursal_id, sucursal.id)

    def test_adjust_stock_creates_new_branch_record_when_missing(self):
        self.client.force_login(self.admin)
        product = create_product('P3', 'Prod 3', precio_compra=Decimal('1200'), precio_venta=Decimal('1800'))
        sucursal = create_sucursal('Sucursal B')

        resp = self.client.post(reverse('ajustar_stock'), {
            'producto_id': product.id,
            'sucursal_id': sucursal.id,
            'delta': '7',
            'motivo': 'Ingreso inicial'
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data.get('success'))
        self.assertEqual(data.get('nueva_cantidad'), 7)

        ss = StockSucursal.objects.get(producto=product, sucursal=sucursal)
        self.assertEqual(ss.cantidad, 7)
        ajuste = AjusteStock.objects.filter(producto=product, sucursal=sucursal).first()
        self.assertIsNotNone(ajuste)
        self.assertEqual(ajuste.cantidad_delta, 7)
        self.assertEqual(ajuste.motivo, 'Ingreso inicial')

    def test_stock_availability_endpoint_returns_branch_quantity(self):
        self.client.force_login(self.admin)
        product = create_product('P4', 'Prod 4', precio_compra=Decimal('1000'), precio_venta=Decimal('1600'))
        sucursal = create_sucursal('Sucursal C')
        StockSucursal.objects.create(producto=product, sucursal=sucursal, cantidad=12)

        resp = self.client.get(reverse('stock_availability'), {
            'producto_id': product.id,
            'sucursal_id': sucursal.id
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data['cantidad'], 12)
        self.assertEqual(data['producto_id'], product.id)
        self.assertEqual(data['sucursal_id'], sucursal.id)

    def test_upload_products_excel_creates_and_updates_products(self):
        self.client.force_login(self.admin)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Productos'
        sheet.append([
            'NOMBRE', 'DESCRIPCION', 'CODIGO 1', 'CODIGO 2', 'CODIGO DE BARRAS',
            'FECHA DE INGRESO', 'PRECIO DE COMPRA', 'PRECIO DE VENTA'
        ])
        sheet.append(['Prod A', 'Desc A', 'X1', 'C2A', 'BAR1', '2026-04-08', 1000, 1500])
        sheet.append(['Prod B', 'Desc B', 'X2', '', 'BAR2', '', 2000, 2500])
        sheet.append(['Prod A v2', 'Desc A2', 'X1', 'C2A2', 'BAR1', '2026-04-08', 1100, 1600])

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        uploaded_file = SimpleUploadedFile(
            'productos.xlsx', buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        resp = self.client.post(reverse('upload_products'), {'file': uploaded_file})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Product.objects.filter(producto_id='X1').count(), 1)
        prod_a = Product.objects.get(producto_id='X1')
        self.assertEqual(prod_a.nombre, 'Prod A v2')
        self.assertEqual(prod_a.descripcion, 'Desc A2')
        self.assertEqual(prod_a.codigo_alternativo, 'C2A2')
        self.assertEqual(prod_a.codigo_barras, 'BAR1')
        self.assertEqual(Product.objects.filter(producto_id='X2').count(), 1)

    def test_download_template_excel_headers_include_optional_labels(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('download_template'))
        self.assertEqual(resp.status_code, 200)
        workbook = load_workbook(filename=BytesIO(resp.content), read_only=True)
        headers = [cell.value for cell in workbook.active[1]]
        self.assertIn('CODIGO 2 (opcional)', headers)
        self.assertIn('CODIGO DE BARRAS (opcional)', headers)
